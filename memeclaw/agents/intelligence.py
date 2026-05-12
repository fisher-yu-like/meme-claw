"""梗情报局 Agent — daily hot meme discovery via search API + LLM judging.

Parallel per-platform pipeline:
  1. Search each platform independently
  2. LLM judges each platform's results in parallel
  3. Merge + deduplicate across platforms
  4. JSON repair with retry (never exposes errors to user)
"""

import asyncio
from datetime import datetime

from memeclaw.config import config
from memeclaw.models.schemas import IntelligenceBriefing, MemeCandidate
from memeclaw.tools.search import search_platform_memes_grouped
from memeclaw.tools.json_repair import JSONRepairAgent


PER_PLATFORM_JUDGE_PROMPT = """你是一个专业的网络热梗鉴定专家，专门评估和筛选中文互联网上的热门网络梗。

你的任务：分析来自**单个平台**的搜索内容片段（可能是单个梗的介绍，也可能是"热梗合集"页面），从中提取和判定网络热梗。

## 重要提示：搜索结果可能是"梗合集"

搜索结果经常是类似"2025年X月热梗合集"的页面摘要。你需要：
- 从合集的摘要中识别出每一个被列举的梗名
- 如果摘要里包含了该梗的简短解释，直接采用
- 如果只有梗名没有解释，根据你对网络文化的了解补充含义
- 对于无法确认含义的梗，标记为 is_hot=false

## 热梗判定标准（三条必须同时满足）：

1. **多人使用**：该梗被至少3个以上不同创作者/账号/帖子引用或讨论，不是单人偶然使用
2. **流传较广**：至少在1个平台上有显著传播（多个相关帖子/视频），或跨平台出现
3. **有明确含义**：能识别出这个梗的大致含义、用法或来源，而不是无意义的随机词组

## 排除标准（出现任一即排除）：

- 仅仅是热门话题/新闻事件，而非网络梗（如社会新闻、政策话题）
- 广告/营销内容伪装成梗
- 无法判断含义的随机词汇
- 已过时超过1个月的旧梗（除非有明显复兴迹象）

## 输出格式：

只输出 JSON，不要附带任何解释文字。严格遵循以下 JSON 语法规则：
- **字符串值内部的双引号必须转义**：用 \\" 而非 "。例如 "他说\\"你好\\"" 而不是 "他说"你好""
- **字符串值内部不得有未转义的换行**：多行文本用 \\n 表示
- platforms 和 urls 用逗号分隔的字符串，不用数组

{
  "candidates": [
    {
      "keyword": "梗的名称",
      "definition": "这个梗的含义和用法（1-2句话，内部引用用\\"转义）",
      "platforms": "douyin",
      "hot_score": 85,
      "is_hot": true,
      "judge_reason": "该词在该平台多个视频标题中出现，含义明确，传播广度足够",
      "urls": "https://example.com/1, https://example.com/2"
    }
  ],
  "summary": "该平台共发现X个候选梗，其中Y个确认为真实热梗。"
}

hot_score 打分规则：
- 90-100：全平台爆火，相关帖子>50条
- 70-89：多平台传播，单一平台热度显著
- 50-69：有一定传播但范围有限
- 30-49：初步出现，可能正在发酵
- 0-29：不符合热梗标准
"""


# ── Per-candidate emergency extraction (used when full JSON is unparsable) ──

def _extract_candidates_individual(text: str) -> tuple[list, str]:
    """Fallback: regex-extract individual candidates from broken JSON.

    When the full JSON is unparsable, this finds each { ... } candidate
    block inside the candidates array and extracts fields via regex.
    Broken candidates are skipped; valid ones are kept.
    """
    import re
    import json as _json

    candidates = []

    arr_start = text.find('"candidates"')
    if arr_start == -1:
        return candidates, ""
    bracket = text.find("[", arr_start)
    if bracket == -1:
        return candidates, ""

    blocks = _split_candidate_blocks(text, bracket + 1)

    for block in blocks:
        try:
            cand = _parse_one_candidate(block, re, _json)
            if cand and cand.keyword:
                candidates.append(cand)
        except Exception:
            pass

    summary = ""
    sm = re.search(r'"summary"\s*:\s*"([^"]*)"', text)
    if sm:
        try:
            summary = _json.loads(f'"{sm.group(1)}"')
        except Exception:
            summary = sm.group(1)

    return candidates, summary


def _split_candidate_blocks(text: str, start: int) -> list[str]:
    """Split the candidates array body into individual { ... } blocks."""
    from memeclaw.tools.json_repair import _escape_control_chars_in_strings, _peek_after_whitespace

    blocks = []
    depth = 0
    block_start = -1
    in_string = False
    escape_next = False
    string_role = ""

    for i in range(start, len(text)):
        ch = text[i]
        if escape_next:
            escape_next = False
            continue
        if ch == "\\" and in_string:
            escape_next = True
            continue
        if ch == '"':
            if not in_string:
                in_string = True
                j = i - 1
                while j >= start and text[j] in " \t\r\n":
                    j -= 1
                string_role = "value" if (j >= start and text[j] == ":") else "key"
            else:
                if string_role == "key":
                    in_string = False
                else:
                    m = _peek_after_whitespace(text, i + 1)
                    if m is None or m in ",:}]":
                        in_string = False
            continue
        if in_string:
            continue
        if ch == "{":
            if depth == 0:
                block_start = i
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0 and block_start != -1:
                blocks.append(text[block_start:i + 1])
                block_start = -1
        elif ch == "]" and depth == 0:
            break

    return blocks


def _parse_one_candidate(block: str, re, _json) -> "MemeCandidate | None":
    """Regex-extract fields from a single candidate JSON block."""
    from memeclaw.tools.json_repair import _escape_control_chars_in_strings

    block = _escape_control_chars_in_strings(block)

    try:
        from json_repair import repair_json
        repaired = repair_json(block, return_objects=False)
        data = _json.loads(repaired)
        return _build_candidate_from_dict(data)
    except Exception:
        pass

    data = {}
    patterns = [
        ("keyword", r'"keyword"\s*:?\s*"((?:[^"\\]|\\.)*)"'),
        ("definition", r'"definition"\s*:?\s*"((?:[^"\\]|\\.)*)"'),
        ("platforms", r'"(?:platforms|source_platforms)"\s*:?\s*"((?:[^"\\]|\\.)*)"'),
        ("judge_reason", r'"judge_reason"\s*:?\s*"((?:[^"\\]|\\.)*)"'),
        ("urls", r'"urls"\s*:?\s*"((?:[^"\\]|\\.)*)"'),
        ("hot_score", r'"hot_score"\s*:\s*(\d+)'),
        ("is_hot", r'"is_hot"\s*:\s*(true|false)'),
    ]

    for key, pattern in patterns:
        m = re.search(pattern, block)
        if m:
            val = m.group(1)
            if key == "hot_score":
                data[key] = int(val)
            elif key == "is_hot":
                data[key] = val == "true"
            elif key in ("platforms", "urls"):
                data[key] = [s.strip() for s in val.split(",") if s.strip()]
            else:
                try:
                    data[key] = _json.loads(f'"{val}"')
                except Exception:
                    data[key] = val

    return _build_candidate_from_dict(data)


def _build_candidate_from_dict(data: dict) -> "MemeCandidate | None":
    """Build a MemeCandidate from a dict, normalizing field formats."""
    keyword = data.get("keyword", "").strip()
    if not keyword:
        return None

    platforms = data.get("platforms", data.get("source_platforms", ""))
    if isinstance(platforms, str):
        platforms = [p.strip() for p in platforms.split(",") if p.strip()]
    elif not isinstance(platforms, list):
        platforms = []

    urls = data.get("urls", [])
    if isinstance(urls, str):
        urls = [u.strip() for u in urls.split(",") if u.strip()]
    elif not isinstance(urls, list):
        urls = []

    return MemeCandidate(
        keyword=keyword,
        definition=data.get("definition", ""),
        source_platforms=platforms,
        hot_score=data.get("hot_score", 0),
        is_hot=data.get("is_hot", False),
        judge_reason=data.get("judge_reason", ""),
        urls=urls,
    )


# ── Intelligence Agent ──

class IntelligenceAgent:
    """Daily surveillance agent for new and trending memes.

    Uses parallel per-platform search + LLM judging instead of
    single-pass merged judging. Each platform is judged independently
    in parallel, then results are merged and deduplicated.

    JSON parsing errors are handled internally with retry — errors
    are never exposed to the user.
    """

    def __init__(self, llm_client=None, db=None):
        self.llm = llm_client
        self.db = db
        self._json_repair = JSONRepairAgent(llm_client=llm_client)

    # ── Main Entry Point ──

    async def run_daily_scan(self, progress_callback=None) -> IntelligenceBriefing:
        """Execute full daily scan: search → parallel per-platform judge → merge.

        If progress_callback is provided, called as:
            await progress_callback(message: str, step: str)
        """
        async def _progress(msg, step):
            if progress_callback:
                await progress_callback(msg, step)

        # Step 1: Search all platforms in parallel, grouped by platform
        await _progress("正在搜索抖音、B站、微博、小红书...", "search")
        try:
            platform_results = await asyncio.wait_for(search_platform_memes_grouped(), timeout=90)
        except asyncio.TimeoutError:
            print("[intelligence] Platform search timed out")
            platform_results = {}

        total_results = sum(len(v) for v in platform_results.values())
        platform_names = list(platform_results.keys())
        print(f"[intelligence] Got {total_results} raw results across {len(platform_names)} platforms: {platform_names}")
        await _progress(
            f"获取到 {total_results} 条搜索结果（{', '.join(platform_names)}），正在并行交由大模型判断...",
            "search_done"
        )

        if not platform_results:
            await _progress("搜索未找到热梗相关内容", "empty")
            return IntelligenceBriefing(
                scan_date=datetime.now(),
                summary="今日搜索未找到任何热梗相关内容，可能是搜索服务暂时不可用。",
            )

        # Step 2: Judge each platform in parallel (each gets its own LLM call)
        await _progress(f"大模型正在并行判断 {len(platform_names)} 个平台的热梗...", "judge")
        judge_tasks = [
            self._judge_platform(platform, results)
            for platform, results in platform_results.items()
            if results
        ]

        judge_results = await asyncio.gather(*judge_tasks, return_exceptions=True)

        # Collect per-platform results
        all_candidates: list[MemeCandidate] = []
        platform_summaries: list[str] = []
        failed_platforms: list[str] = []

        for platform, result in zip(
            [p for p in platform_results if platform_results[p]], judge_results
        ):
            if isinstance(result, Exception):
                print(f"[intelligence] Platform '{platform}' judge failed: {result}")
                failed_platforms.append(platform)
                continue
            candidates, summary = result
            all_candidates.extend(candidates)
            if summary:
                platform_summaries.append(f"[{platform}] {summary}")

        # Step 3: Deduplicate across platforms
        all_candidates = _deduplicate_candidates(all_candidates)

        # Step 4: Separate confirmed hot memes
        confirmed = [c for c in all_candidates if c.is_hot and c.hot_score >= config.hot_score_threshold]

        # Build merged summary
        summary = _merge_summaries(platform_summaries, all_candidates, confirmed, failed_platforms)

        await _progress(f"判断完成：{len(all_candidates)} 个候选，{len(confirmed)} 个确认为热梗，正在生成简报...", "format")
        print(f"[intelligence] Found {len(all_candidates)} candidates, {len(confirmed)} confirmed hot across {len(platform_names)} platforms")

        return IntelligenceBriefing(
            scan_date=datetime.now(),
            candidates=all_candidates,
            confirmed_hot_memes=confirmed,
            summary=summary,
        )

    # ── Per-Platform Judging ──

    async def _judge_platform(
        self, platform: str, search_results: list[dict]
    ) -> tuple[list[MemeCandidate], str]:
        """Judge a single platform's search results with the LLM.

        Called in parallel for each platform. Uses JSON repair agent
        with retry — errors are never raised to the caller.
        """
        if self.llm is None:
            return _mock_judge_platform(platform, search_results)

        context = self._format_search_context(platform, search_results)

        # Try LLM call + JSON extraction with internal retry
        for attempt in range(3):  # up to 3 attempts per platform
            try:
                response = await self.llm.chat.completions.create(
                    model=config.long_context_model,
                    max_tokens=4096,
                    timeout=600.0,
                    messages=[
                        {"role": "system", "content": PER_PLATFORM_JUDGE_PROMPT},
                        {"role": "user", "content": f"以下是 **{platform}** 平台今日搜索到的可能热梗内容，请逐一判断：\n\n{context}"},
                    ],
                )
                text = response.choices[0].message.content
                if not text or not text.strip():
                    reasoning = getattr(response.choices[0].message, "reasoning_content", None)
                    if reasoning:
                        text = reasoning
                    else:
                        continue  # retry

                # Use JSON repair agent (internal retry with LLM re-prompt)
                data = await self._json_repair.extract_with_retry(text, max_retries=1)
                if not data:
                    # Try per-candidate emergency extraction
                    candidates, summary = _extract_candidates_individual(text)
                    if candidates:
                        print(f"[intelligence] [{platform}] Per-candidate extraction rescued {len(candidates)} candidates")
                        return candidates, summary
                    continue  # retry the whole LLM call

                candidates = self._parse_candidates_from_data(data)
                summary = data.get("summary", "")
                print(f"[intelligence] [{platform}] Judge found {len(candidates)} candidates")
                return candidates, summary

            except Exception as e:
                if attempt < 2:
                    print(f"[intelligence] [{platform}] Judge attempt {attempt + 1} failed: {e}, retrying...")
                    await asyncio.sleep(1)  # brief pause before retry
                else:
                    print(f"[intelligence] [{platform}] Judge failed after 3 attempts: {e}")

        # All retries exhausted — return empty, caller handles gracefully
        print(f"[intelligence] [{platform}] All judge attempts exhausted, returning empty")
        return [], ""

    def _parse_candidates_from_data(self, data: dict) -> list[MemeCandidate]:
        """Parse MemeCandidate list from extracted JSON data."""
        candidates = []
        for item in data.get("candidates", []):
            platforms = item.get("platforms", item.get("source_platforms", ""))
            if isinstance(platforms, str):
                platforms = [p.strip() for p in platforms.split(",") if p.strip()]
            elif not isinstance(platforms, list):
                platforms = []

            urls = item.get("urls", [])
            if isinstance(urls, str):
                urls = [u.strip() for u in urls.split(",") if u.strip()]
            elif not isinstance(urls, list):
                urls = []

            candidates.append(MemeCandidate(
                keyword=item.get("keyword", ""),
                definition=item.get("definition", ""),
                source_platforms=platforms,
                hot_score=item.get("hot_score", 0),
                is_hot=item.get("is_hot", False),
                judge_reason=item.get("judge_reason", ""),
                urls=urls,
            ))
        return candidates

    def _format_search_context(self, platform: str, results: list[dict]) -> str:
        """Format one platform's search results into LLM-readable context."""
        lines = [f"### {platform}"]
        for i, item in enumerate(results[:8], 1):
            lines.append(
                f"[{i}] {item.get('title', '')}\n"
                f"    摘要: {item.get('snippet', '')[:100]}\n"
                f"    链接: {item.get('url', '')}"
            )
        return "\n".join(lines)

    # ── Formatting & Export ──

    def format_briefing(self, briefing: IntelligenceBriefing) -> str:
        """Render a daily briefing as a user-friendly Markdown message."""
        lines = [
            f"## 梗情报局日报 — {briefing.scan_date.strftime('%Y-%m-%d')}",
            "",
            briefing.summary,
        ]
        if briefing.confirmed_hot_memes:
            lines.append("\n### 🔥 今日确认热梗")
            for i, m in enumerate(briefing.confirmed_hot_memes[:15], 1):
                platforms = "、".join(m.source_platforms) if m.source_platforms else "?"
                lines.append(f"{i}. **{m.keyword}**  [{platforms}] 热度{m.hot_score}")
                if m.definition:
                    lines.append(f"   > {m.definition}")
                if m.judge_reason:
                    lines.append(f"   - 判定：{m.judge_reason}")
        elif briefing.candidates:
            lines.append("\n### 📋 待确认候选")
            for i, m in enumerate(briefing.candidates[:10], 1):
                lines.append(f"{i}. **{m.keyword}** — 热度{m.hot_score}（{m.judge_reason}）")
        else:
            lines.append("\n今日暂未发现新的热梗，请明日再查看。")
        return "\n".join(lines)

    def export_briefing(self, briefing: IntelligenceBriefing) -> str:
        """Export daily briefing to local files (md + xlsx). Cleans up files older than 7 days."""
        import os
        import glob
        import time

        export_dir = config.export_dir
        os.makedirs(export_dir, exist_ok=True)

        date_str = briefing.scan_date.strftime("%Y-%m-%d")
        base = os.path.join(export_dir, f"daily_briefing_{date_str}")

        md_content = self.format_briefing(briefing)
        md_path = base + ".md"
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)

        xlsx_path = base + ".xlsx"
        rows = self._build_export_rows(briefing)
        try:
            self._write_xlsx(xlsx_path, rows)
        except Exception:
            csv_path = base + ".csv"
            self._write_csv(csv_path, rows)

        cutoff = time.time() - 7 * 86400
        for old in glob.glob(os.path.join(export_dir, "daily_briefing_*")):
            if os.path.getmtime(old) < cutoff:
                try:
                    os.remove(old)
                except OSError:
                    pass

        return md_path

    def _build_export_rows(self, briefing: IntelligenceBriefing) -> list[dict]:
        rows = []
        source = briefing.confirmed_hot_memes if briefing.confirmed_hot_memes else briefing.candidates
        for i, m in enumerate(source[:20], 1):
            rows.append({
                "序号": i,
                "梗名": m.keyword,
                "含义": m.definition,
                "出现平台": "、".join(m.source_platforms),
                "热度评分": m.hot_score,
                "判定理由": m.judge_reason,
                "状态": "🔥确认热梗" if m.is_hot else "待确认",
            })
        return rows

    def _write_xlsx(self, path: str, rows: list[dict]) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "每日热梗推送"

        headers = ["序号", "梗名", "含义", "出现平台", "热度评分", "判定理由", "状态"]
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
        header_alignment = Alignment(horizontal="center")
        hot_fill = PatternFill(start_color="FFE4E1", end_color="FFE4E1", fill_type="solid")

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        for row_idx, row in enumerate(rows, 2):
            for col_idx, key in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
                if row.get("状态") == "🔥确认热梗":
                    cell.fill = hot_fill

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 14
        ws.column_dimensions["E"].width = 10
        ws.column_dimensions["F"].width = 40
        ws.column_dimensions["G"].width = 14

        wb.save(path)

    def _write_csv(self, path: str, rows: list[dict]) -> None:
        import csv
        headers = ["序号", "梗名", "含义", "出现平台", "热度评分", "判定理由", "状态"]
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)


# ── Module-level helpers ──

def _deduplicate_candidates(candidates: list[MemeCandidate]) -> list[MemeCandidate]:
    """Deduplicate candidates across platforms by keyword similarity.

    When the same meme appears on multiple platforms, merge their
    platform lists, URLs, and take the best definition + score.
    """
    if not candidates:
        return []

    # Group by normalized keyword
    groups: dict[str, list[MemeCandidate]] = {}
    for c in candidates:
        key = c.keyword.strip().lower()
        # Also try to match by common variations
        found_key = None
        for existing in groups:
            if key == existing or key in existing or existing in key:
                found_key = existing
                break
        if found_key:
            groups[found_key].append(c)
        else:
            groups[key] = [c]

    merged = []
    for group in groups.values():
        if len(group) == 1:
            merged.append(group[0])
            continue

        # Merge: best definition (longest non-empty), union platforms/urls,
        # max hot_score, OR of is_hot, combine judge_reasons
        best_def = ""
        all_platforms: set[str] = set()
        all_urls: set[str] = set()
        best_score = 0
        any_hot = False
        reasons: list[str] = []

        for c in group:
            if len(c.definition) > len(best_def):
                best_def = c.definition
            all_platforms.update(c.source_platforms)
            all_urls.update(c.urls)
            if c.hot_score > best_score:
                best_score = c.hot_score
            if c.is_hot:
                any_hot = True
            if c.judge_reason:
                reasons.append(c.judge_reason)

        merged.append(MemeCandidate(
            keyword=group[0].keyword,
            definition=best_def,
            source_platforms=sorted(all_platforms),
            hot_score=best_score,
            is_hot=any_hot,
            judge_reason=" | ".join(reasons),
            urls=sorted(all_urls),
        ))

    # Sort by hot_score descending
    merged.sort(key=lambda c: c.hot_score, reverse=True)
    return merged


def _merge_summaries(
    platform_summaries: list[str],
    all_candidates: list[MemeCandidate],
    confirmed: list[MemeCandidate],
    failed_platforms: list[str],
) -> str:
    """Merge per-platform summaries into a single briefing summary."""
    parts = []

    parts.append(f"今日共扫描 {len(all_candidates)} 个候选梗，确认 {len(confirmed)} 个为真实热梗。")

    if confirmed:
        names = "、".join(c.keyword for c in confirmed[:5])
        if len(confirmed) > 5:
            names += "等"
        parts.append(f"主要热点：{names}。")

    if failed_platforms:
        parts.append(f"（{'、'.join(failed_platforms)} 平台暂时无法获取数据）")

    return " ".join(parts)


def _mock_judge_platform(platform: str, search_results: list[dict]) -> tuple[list[MemeCandidate], str]:
    """Fallback when LLM is unavailable — basic heuristic extraction for one platform."""
    candidates = []
    seen = set()
    for r in search_results[:10]:
        title = r.get("title", "")
        if title and title not in seen:
            seen.add(title)
            candidates.append(MemeCandidate(
                keyword=title[:30],
                definition="",
                source_platforms=[platform],
                hot_score=50,
                is_hot=False,
                judge_reason="LLM 不可用，基于搜索结果自动提取",
                urls=[r.get("url", "")] if r.get("url") else [],
            ))
    return candidates, f"[{platform}] LLM 判断不可用，以下为原始搜索结果，需人工筛选。"
